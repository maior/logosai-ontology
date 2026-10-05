"""
readers 포맷 확장 — docx / hwp / doc / jsonld.

업로드는 pdf·csv·json 만 오지 않는다 (사용자 요구: hwp·doc 포함, 분야 다양).
원칙은 기존 readers 와 동일하다: 포맷 변환은 **결정적, LLM 무사용**. LLM 은
내용 판단(감식·추출)에만 쓴다 — 바이트를 텍스트로 바꾸는 데 판단은 없다.

HWP 는 실파일을 만들 수 없으므로(olefile 은 읽기 전용) 레코드 파서를 순수
함수로 뽑아 합성 바이트로 검증한다 — 포맷 명세(HWP 5.0 레코드 헤더)가 계약이다.
"""

import struct

import pytest

from ontology.builder.readers import (
    SUPPORTED_EXTENSIONS,
    UnsupportedFormatError,
    _decode_hwp_text,
    _parse_hwp_section,
    read_file,
)


class TestSupportedExtensions:
    def test_new_formats_are_registered(self):
        for ext in (".docx", ".doc", ".hwp", ".jsonld"):
            assert ext in SUPPORTED_EXTENSIONS, ext

    def test_unknown_still_raises(self, tmp_path):
        p = tmp_path / "x.xyz"
        p.write_text("data")
        with pytest.raises(UnsupportedFormatError):
            read_file(p)


class TestDocx:
    def test_paragraphs_and_tables_are_extracted(self, tmp_path):
        import docx

        doc = docx.Document()
        doc.add_paragraph("제1조 청약철회")
        doc.add_paragraph("계약자는 15일 이내에 철회할 수 있다.")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "구분"
        table.rows[0].cells[1].text = "금액"
        table.rows[1].cells[0].text = "보험료"
        table.rows[1].cells[1].text = "10000"
        path = tmp_path / "약관.docx"
        doc.save(str(path))

        text = read_file(path)
        assert "제1조 청약철회" in text
        assert "구분 | 금액" in text        # 표는 행 단위로 편다
        assert "보험료 | 10000" in text


class TestHwpRecordParser:
    """HWP 5.0 레코드 파서 — 합성 바이트로 명세 검증.

    레코드 헤더 = 4바이트 LE DWORD: tag(10bit) | level(10bit) | size(12bit).
    HWPTAG_PARA_TEXT = 67. 본문은 UTF-16LE, 제어문자 32 미만 중 확장 컨트롤은
    16바이트를 차지한다.
    """

    @staticmethod
    def record(tag: int, payload: bytes) -> bytes:
        header = tag | (0 << 10) | (len(payload) << 20)
        return struct.pack("<I", header) + payload

    def test_para_text_is_extracted(self):
        payload = "안녕하세요".encode("utf-16-le")
        data = self.record(67, payload)
        assert _parse_hwp_section(data) == "안녕하세요"

    def test_non_text_records_are_skipped(self):
        data = (self.record(66, b"\x00" * 8)          # 다른 태그 — 무시
                + self.record(67, "본문".encode("utf-16-le")))
        assert _parse_hwp_section(data) == "본문"

    def test_extended_control_is_skipped_entirely(self):
        """확장 컨트롤(표·그림 등)은 WCHAR 8개 = 16바이트를 통째로 차지한다.
        1개만 걷어내면 나머지 7개가 쓰레기 글자로 새어 나온다."""
        payload = ("앞".encode("utf-16-le")
                   + struct.pack("<H", 11)            # 확장 컨트롤 시작
                   + b"\x41\x00" * 7                  # 컨트롤 내부 데이터 ("A"×7)
                   + "뒤".encode("utf-16-le"))
        assert _decode_hwp_text(payload) == "앞뒤"

    def test_line_break_controls_become_newline(self):
        payload = ("가".encode("utf-16-le") + struct.pack("<H", 13)
                   + "나".encode("utf-16-le"))
        assert _decode_hwp_text(payload) == "가\n나"

    def test_oversized_record_header_is_handled(self):
        """size 필드가 0xFFF 면 실제 크기는 다음 4바이트에 온다 (명세)."""
        payload = "긴본문".encode("utf-16-le")
        header = 67 | (0xFFF << 20)
        data = struct.pack("<I", header) + struct.pack("<I", len(payload)) + payload
        assert _parse_hwp_section(data) == "긴본문"

    def test_truncated_data_does_not_raise(self):
        """잘린 파일 — 죽지 말고 읽은 데까지 돌려준다."""
        data = self.record(67, "본문".encode("utf-16-le")) + b"\x43"
        assert _parse_hwp_section(data) == "본문"

    def test_not_an_hwp_file_raises_clear_error(self, tmp_path):
        p = tmp_path / "fake.hwp"
        p.write_bytes(b"this is not an ole file")
        with pytest.raises(Exception):   # olefile 거부 또는 명시 에러
            read_file(p)


class TestDoc:
    def test_doc_error_names_the_remedy(self, tmp_path, monkeypatch):
        """textutil(macOS) 이 없으면, 뭘 하라는지 말하는 에러여야 한다."""
        import shutil as _shutil

        monkeypatch.setattr(_shutil, "which", lambda cmd: None)
        p = tmp_path / "old.doc"
        p.write_bytes(b"\xd0\xcf\x11\xe0old doc bytes")
        with pytest.raises(UnsupportedFormatError, match="docx"):
            read_file(p)
